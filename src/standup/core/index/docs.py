"""How much the project's own writing dwells on each file. Semantic signal that costs nothing."""

import hashlib
import re
import zipfile
from collections import Counter
from contextlib import contextmanager
from pathlib import Path

import tree_sitter_language_pack as tsl
from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageError
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from standup.core.index.code import walk
from standup.core.llm import describe_image_sync
from standup.core.models import FileFacts
from standup.errors import InvalidInput

DOC_SUFFIXES = {".md", ".rst", ".txt", ".adoc"}
PDF_SUFFIX = ".pdf"
OFFICE_SUFFIXES = {".xlsx", ".pptx", ".docx"}
IMAGE_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
ATTACHABLE = DOC_SUFFIXES | {PDF_SUFFIX} | OFFICE_SUFFIXES | set(IMAGE_MEDIA_TYPES)
MAX_DOC_CHARS = 4_000_000
MIN_NAME_LENGTH = 5
_WORD = re.compile(r"[\w?!-]+")

IMAGE_PROMPT = (
    "Describe this image as evidence for a status update: what it shows, any text or data "
    "visible in it verbatim, and what it demonstrates about the underlying work."
)

_DESCRIBED: dict[str, str] = {}
_DESCRIBED_MAX = 16


def readable(name: str) -> bool:
    """Whether a file attached on its own can be turned into text. Checked before it is kept."""
    suffix = Path(name).suffix.lower()
    return suffix in ATTACHABLE or tsl.detect_language_from_path(name) is not None


def read(path: Path) -> str:
    """A file as text. The one place that knows how to open a kind of document."""
    suffix = path.suffix.lower()
    if suffix == PDF_SUFFIX:
        return _pdf(path)
    if suffix == ".xlsx":
        return _xlsx(path)
    if suffix == ".pptx":
        return _pptx(path)
    if suffix == ".docx":
        return _docx(path)
    if suffix in IMAGE_MEDIA_TYPES:
        return _image(path, suffix)
    with _reading(path.name, OSError, UnicodeDecodeError):
        return path.read_text(encoding="utf-8")


@contextmanager
def _reading(name: str, *errors: type[Exception]):
    """Every format-specific reader raises the same way; only the library errors differ. Wraps
    the whole read, not just opening the file — a truncated body can fail as easily as a bad
    header, and a stream reader like openpyxl's `read_only` mode may not notice until it iterates."""
    try:
        yield
    except errors as unreadable:
        raise InvalidInput(f"{name} could not be read: {unreadable}") from unreadable


def _pdf(path: Path) -> str:
    with _reading(path.name, OSError, PyPdfError, ValueError):
        return "\n".join(page.extract_text() for page in PdfReader(path).pages)


def _xlsx(path: Path) -> str:
    with _reading(path.name, OSError, zipfile.BadZipFile, InvalidFileException):
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = (row for sheet in book.worksheets for row in sheet.iter_rows(values_only=True))
            return "\n".join("\t".join(str(cell) for cell in row if cell is not None) for row in rows)
        finally:
            book.close()


def _pptx(path: Path) -> str:
    with _reading(path.name, OSError, zipfile.BadZipFile, PptxPackageError):
        deck = Presentation(path)
        frames = (s.text_frame for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        cells = (
            cell.text_frame
            for slide in deck.slides
            for shape in slide.shapes
            if shape.has_table
            for row in shape.table.rows
            for cell in row.cells
        )
        return "\n".join(frame.text for frame in (*frames, *cells) if frame.text)


def _docx(path: Path) -> str:
    with _reading(path.name, OSError, zipfile.BadZipFile, DocxPackageError):
        document = Document(path)
        cells = (cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(text for text in (*(p.text for p in document.paragraphs), *cells) if text)


def _image(path: Path, suffix: str) -> str:
    with _reading(path.name, OSError):
        data = path.read_bytes()

    digest = hashlib.sha256(data).hexdigest()
    if digest not in _DESCRIBED:
        if len(_DESCRIBED) >= _DESCRIBED_MAX:
            _DESCRIBED.pop(next(iter(_DESCRIBED)))
        _DESCRIBED[digest] = describe_image_sync(data, IMAGE_MEDIA_TYPES[suffix], prompt=IMAGE_PROMPT)
    return _DESCRIBED[digest]


def prose(root: Path) -> str:
    """Everything the project writes about itself. Only prose files — a PDF buried in a
    repository is a binary nobody asked us to open; one attached on its own is read."""
    collected: list[str] = []
    size = 0
    for path in walk(root):
        if path.suffix.lower() not in DOC_SUFFIXES or size >= MAX_DOC_CHARS:
            continue
        try:
            text = read(path)
        except InvalidInput:
            continue
        collected.append(text)
        size += len(text)
    return "\n".join(collected)


def emphasis(text: str, files: list[FileFacts]) -> dict[str, float]:
    if not text:
        return {}
    lowered = text.lower()
    mentioned = Counter(_WORD.findall(text))

    rates = {}
    for facts in files:
        filename = Path(facts.path).name
        names = [s.name for s in facts.symbols if len(s.name) >= MIN_NAME_LENGTH]
        mentions = sum(mentioned[name] for name in names)
        if len(filename) >= MIN_NAME_LENGTH:
            names.append(filename)
            mentions += lowered.count(filename.lower())
        if mentions:
            rates[facts.path] = mentions / len(names)

    if not rates:
        return {}
    loudest = max(rates.values())
    return {path: rate / loudest for path, rate in rates.items()}
