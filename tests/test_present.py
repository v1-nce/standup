import struct
import zlib
from datetime import UTC, datetime
from pathlib import Path

import pytest
from pptx import Presentation

from standup.core.models import (
    Candidate,
    Commit,
    DeckDesign,
    FileFacts,
    Index,
    Scope,
    Scored,
    Selection,
    Slide,
    SlidePlan,
    Symbol,
    VisualElement,
)
from standup.core.present import build, evidence, order_fault, revise
from standup.core.present.validate import problems, slide_problems
from standup.errors import InvalidInput
from tests.conftest import TINY_PNG

TODAY = datetime(2026, 8, 3, tzinfo=UTC)


@pytest.fixture
def index():
    return Index(
        fingerprint="f",
        built_at=TODAY,
        files=[
            FileFacts(
                path="src/auth.py",
                content_hash="a",
                symbols=[Symbol(name="login", kind="Function", line=1)],
            ),
            FileFacts(path="src/billing.py", content_hash="b"),
        ],
        commits=[
            Commit(
                sha="abc",
                authored_at=TODAY,
                author="vince",
                message="signed tokens replace the session store",
                changes={"src/auth.py": 52},
            )
        ],
    )


@pytest.fixture
def selection():
    return Selection(
        request="standup tomorrow, the auth thing",
        scope=Scope(keywords=["auth"], audience="the team", slide_budget=2),
        chosen=[
            Scored(
                candidate=Candidate(
                    id="src/auth.py", title="auth.py", paths=["src/auth.py"], commits=["abc"]
                ),
                signals={"churn": 1.0},
                score=2.4,
            ),
            Scored(
                candidate=Candidate(
                    id="src/billing.py", title="billing.py", paths=["src/billing.py"]
                ),
                signals={"churn": 0.1},
                score=0.3,
            ),
        ],
        cut=[
            Scored(
                candidate=Candidate(id="docs/notes.md", title="notes.md"),
                signals={},
                score=0.01,
            )
        ],
    )


def good_plan() -> SlidePlan:
    return SlidePlan(
        slides=[
            Slide(candidate_id="src/auth.py", title="Auth", bullets=["src/auth.py now signs tokens"]),
            Slide(candidate_id="src/billing.py", title="Billing", bullets=["Untouched this week"]),
        ]
    )


def test_a_clean_plan_has_nothing_wrong_with_it(selection, index):
    assert problems(good_plan(), selection, index) == []


def test_order_fault_is_none_exactly_when_the_slides_match_selection_order(selection):
    assert order_fault(good_plan().slides, selection) is None

    wrong_order = list(reversed(good_plan().slides))
    fault = order_fault(wrong_order, selection)
    assert fault is not None and "in that order" in fault


@pytest.mark.parametrize(
    ("slides", "fault"),
    [
        ([], "must be exactly"),
        ([Slide(candidate_id="src/auth.py", title="Only one")], "must be exactly"),
        (
            [
                Slide(candidate_id="src/billing.py", title="Billing"),
                Slide(candidate_id="src/auth.py", title="Auth"),
            ],
            "in that order",
        ),
        (
            [
                Slide(candidate_id="src/auth.py", title="Auth"),
                Slide(candidate_id="src/auth.py", title="Auth again"),
            ],
            "must be exactly",
        ),
    ],
)
def test_the_plan_cannot_add_drop_reorder_or_repeat_a_slide(selection, index, slides, fault):
    assert any(fault in problem for problem in problems(SlidePlan(slides=slides), selection, index))


def test_a_file_that_does_not_exist_is_refused(selection, index):
    drafted = good_plan()
    drafted.slides[0].bullets = ["Rewrote src/auth/oauth.py to issue tokens"]
    assert any("no file 'src/auth/oauth.py'" in p for p in problems(drafted, selection, index))


def test_a_symbol_that_does_not_exist_is_refused(selection, index):
    drafted = good_plan()
    drafted.slides[0].bullets = ["refresh_token() was added"]
    assert any("no 'refresh_token'" in p for p in problems(drafted, selection, index))


def test_prose_that_merely_looks_like_a_path_is_left_alone(selection, index):
    drafted = good_plan()
    drafted.slides[0].bullets = ["Tokens are signed and/or encrypted, e.g. on login"]
    assert problems(drafted, selection, index) == []


def test_a_symbol_that_does_not_exist_is_refused_even_called_with_arguments(selection, index):
    """`refresh_token()` was caught; `refresh_token(user)` used to sail through unchecked - the
    more specific claim, naming what's passed in, is the one worth checking."""
    drafted = good_plan()
    drafted.slides[0].bullets = ["refresh_token(user) was added"]
    assert any("no 'refresh_token'" in p for p in problems(drafted, selection, index))


def test_a_real_symbol_called_with_arguments_is_accepted(selection, index):
    drafted = good_plan()
    drafted.slides[0].bullets = ["login(request) now issues a signed token"]
    assert problems(drafted, selection, index) == []


def test_a_parenthetical_is_not_misread_as_a_call(selection, index):
    """A space before the paren means prose, not code - "Figure (below)" is not `Figure(below)`."""
    drafted = good_plan()
    drafted.slides[0].bullets = ["Tokens are signed (see the diagram below)"]
    assert problems(drafted, selection, index) == []


def test_evidence_carries_the_commits_but_never_the_scores_or_the_cut(selection, index):
    cited = evidence(selection, index)

    assert "signed tokens replace the session store" in cited
    assert "src/auth.py" in cited and "src/billing.py" in cited
    assert "docs/notes.md" not in cited
    for banned in ("2.4", "0.3", "churn"):
        assert banned not in cited


def test_evidence_carries_the_symbols_and_the_source_a_specific_bullet_needs(selection, index):
    """Without these a slide can only paraphrase a commit message, which is every deck this
    produced. `defines` also closes the gap where validation rejected calls it never showed."""
    cited = evidence(selection, index, {"src/auth.py": "def login(user):\n    return sign(user)"})

    assert "defines: login" in cited
    assert "def login(user)" in cited


def test_evidence_still_stands_up_when_nothing_can_be_read_from_disk(selection, index):
    cited = evidence(selection, index)
    assert "src/auth.py" in cited
    assert "source:" not in cited


def test_a_dotfile_is_recognised_rather_than_misread_as_fabricated():
    """token.strip("./") used to eat the leading dot off a real path like .github/....yml."""
    dotted = Index(
        fingerprint="f", built_at=TODAY, files=[FileFacts(path=".github/workflows/ci.yml", content_hash="c")]
    )
    scoped = Selection(
        request="ci",
        scope=Scope(slide_budget=1),
        chosen=[
            Scored(
                candidate=Candidate(
                    id=".github/workflows/ci.yml", title="ci.yml", paths=[".github/workflows/ci.yml"]
                ),
                signals={},
                score=1.0,
            )
        ],
    )
    plan = SlidePlan(
        slides=[
            Slide(candidate_id=".github/workflows/ci.yml", title="CI", bullets=["Updated .github/workflows/ci.yml"])
        ]
    )
    assert problems(plan, scoped, dotted) == []


def test_a_slide_may_cite_a_real_file_that_belongs_to_another_candidate(selection, index):
    """Was refused. Scoping names per-candidate rejected true statements, so the only bullets that
    survived were the ones specific about nothing - which is what made every deck read the same."""
    drafted = good_plan()
    drafted.slides[1].bullets = ["Related work landed in src/auth.py"]
    assert problems(drafted, selection, index) == []


def test_a_method_call_on_someone_elses_object_is_not_read_as_a_project_symbol(selection, index):
    """`deck.save()` is python-pptx's, not ours; only a bare call claims something this project has."""
    drafted = good_plan()
    drafted.slides[0].bullets = ["The renderer now calls deck.save() once per download"]
    assert problems(drafted, selection, index) == []


def test_a_deleted_file_can_still_be_written_about(selection, index):
    index.files = [facts for facts in index.files if facts.path != "src/auth.py"]
    drafted = good_plan()
    drafted.slides[0].bullets = ["src/auth.py was removed"]
    assert problems(drafted, selection, index) == []


def test_removing_and_reordering_a_selection_costs_no_model_call(selection):
    selection.chosen.reverse()
    order = [entry.candidate.id for entry in selection.chosen]
    revised = revise(good_plan(), order)
    assert [s.candidate_id for s in revised.slides] == ["src/billing.py", "src/auth.py"]

    selection.chosen.pop()
    order = [entry.candidate.id for entry in selection.chosen]
    assert [s.candidate_id for s in revise(good_plan(), order).slides] == ["src/billing.py"]


def test_restoring_a_cut_item_needs_a_slide_written_for_it(selection):
    selection.chosen.append(selection.cut[0])
    order = [entry.candidate.id for entry in selection.chosen]
    with pytest.raises(InvalidInput, match="docs/notes.md"):
        revise(good_plan(), order)


def test_revise_interleaves_a_free_slide_wherever_the_order_places_it(selection):
    """Free ids ride along with candidate ids in one order - `revise` doesn't treat them specially."""
    plan = good_plan()
    plan.slides.insert(0, Slide(candidate_id="title", free=True, title="Standup"))
    order = ["title", "src/billing.py", "src/auth.py"]

    revised = revise(plan, order)
    assert [s.candidate_id for s in revised.slides] == order


def test_revise_rejects_a_free_id_with_no_slide_written_yet(selection):
    order = [entry.candidate.id for entry in selection.chosen] + ["title"]
    with pytest.raises(InvalidInput, match="title"):
        revise(good_plan(), order)


def test_a_free_slide_is_never_checked_against_the_evidence(selection, index):
    """No candidate backs a free slide, so nothing about it can be fabricated in the first place."""
    free = Slide(candidate_id="title", free=True, title="FINNATO", bullets=["src/nonexistent.py"])
    assert slide_problems([free], index) == []


def test_a_free_slide_does_not_break_the_completeness_check(selection, index):
    plan = good_plan()
    plan.slides.insert(0, Slide(candidate_id="title", free=True, title="Standup"))
    assert problems(plan, selection, index) == []


def test_a_free_slide_can_cite_an_attached_image(selection, index):
    index.files.append(FileFacts(path="doc/baitrage.png", content_hash="p"))
    free = Slide(candidate_id="conclusion", free=True, title="Conclusion", image="doc/baitrage.png")
    assert slide_problems([free], index) == []


def test_an_image_that_is_not_attached_is_refused(selection, index):
    free = Slide(candidate_id="conclusion", free=True, title="Conclusion", image="doc/nope.png")
    assert any("no attached image" in p for p in slide_problems([free], index))


def test_an_image_pointing_at_a_non_image_file_is_refused(selection, index):
    """A real indexed path, but not a picture - src/auth.py exists, .py is not a picture format."""
    free = Slide(candidate_id="conclusion", free=True, title="Conclusion", image="src/auth.py")
    assert any("no attached image" in p for p in slide_problems([free], index))


def _canvas_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def test_an_evidence_slide_can_use_a_grounded_attached_image(selection, index):
    index.files.append(FileFacts(path="doc/baitrage.png", content_hash="p"))
    drafted = good_plan()
    drafted.slides[0].image = "doc/baitrage.png"
    assert slide_problems(drafted.slides, index) == []


def test_the_deck_opens_with_the_slides_it_was_given(tmp_path):
    written = build(good_plan(), tmp_path / "decks" / "standup.pptx")
    assert written.is_file()

    reopened = Presentation(str(written))
    assert len(reopened.slides) == 2
    assert "Auth" in _canvas_text(reopened.slides[0])
    assert "Billing" in _canvas_text(reopened.slides[1])
    assert "src/auth.py now signs tokens" in _canvas_text(reopened.slides[0])


def test_a_slide_with_several_bullets_keeps_them_all(tmp_path):
    many = SlidePlan(
        slides=[Slide(candidate_id="src/auth.py", title="Auth", bullets=["one", "two", "three"])]
    )
    written = build(many, Path(tmp_path) / "many.pptx")
    canvas = _canvas_text(Presentation(str(written)).slides[0])
    assert all(item in canvas for item in ("one", "two", "three"))


def test_a_bare_free_slide_gets_a_cover_canvas_not_a_bullet_template(tmp_path):
    plan = SlidePlan(slides=[Slide(candidate_id="cover", free=True, title="ArbiSuits")])
    written = build(plan, tmp_path / "cover.pptx")
    slide = Presentation(str(written)).slides[0]
    assert "ArbiSuits" in _canvas_text(slide)
    assert slide.slide_layout.name == "Blank"
    assert "•" not in _canvas_text(slide)


def test_semantic_layouts_render_different_topologies(tmp_path):
    plan = SlidePlan(
        design=DeckDesign(theme="editorial"),
        slides=[
            Slide(candidate_id="section", free=True, title="Decisions", layout="section"),
            Slide(
                candidate_id="comparison",
                free=True,
                title="Before and after",
                bullets=["Session lookup"],
                secondary_title="After",
                secondary_bullets=["Signed token"],
                layout="two_column",
            ),
            Slide(candidate_id="takeaway", free=True, title="One path now owns auth", layout="statement"),
        ],
    )
    reopened = Presentation(str(build(plan, tmp_path / "layouts.pptx")))

    assert str(reopened.slides[0].background.fill.fore_color.rgb) == "B54432"
    assert "Session lookup" in _canvas_text(reopened.slides[1])
    assert "Signed token" in _canvas_text(reopened.slides[1])
    assert len(reopened.slides[1].shapes) != len(reopened.slides[2].shapes)


def test_a_custom_canvas_renders_ordered_editable_primitives(tmp_path):
    plan = SlidePlan(
        design=DeckDesign(theme="dark", accent="#23D5AB"),
        slides=[
            Slide(
                candidate_id="canvas",
                free=True,
                title="Canvas",
                elements=[
                    VisualElement(
                        id="panel", kind="shape", shape="rounded", x=5, y=8, width=90, height=84,
                        fill="surface",
                    ),
                    VisualElement(
                        id="rule", kind="line", x=9, y=23, width=28, height=0,
                        stroke="accent", stroke_width=3,
                    ),
                    VisualElement(
                        id="headline", kind="text", x=9, y=30, width=74, height=28,
                        text="A designed argument", color="text", font_size=44, font_weight="bold",
                    ),
                    VisualElement(
                        id="orb", kind="shape", shape="ellipse", x=76, y=60, width=14, height=25,
                        fill="#23D5AB", rotation=12,
                    ),
                ],
            )
        ],
    )

    slide = Presentation(str(build(plan, tmp_path / "canvas.pptx"))).slides[0]
    assert "A designed argument" in _canvas_text(slide)
    assert len(slide.shapes) == 4
    assert str(slide.background.fill.fore_color.rgb) == "0B1020"


def test_canvas_geometry_cannot_escape_the_slide():
    with pytest.raises(ValueError, match="inside the 100×100 canvas"):
        VisualElement(id="outside", kind="shape", x=90, y=10, width=20, height=20)


def test_canvas_text_on_an_evidence_slide_still_crosses_the_grounding_boundary(index):
    slide = Slide(
        candidate_id="src/auth.py",
        title="Auth",
        elements=[
            VisualElement(
                id="claim", kind="text", x=5, y=5, width=80, height=20,
                text="Replaced src/ghost.py",
            )
        ],
    )
    assert any("src/ghost.py" in problem for problem in slide_problems([slide], index))


def test_an_image_slide_embeds_a_real_picture_not_literal_markdown(tmp_path):
    photo = tmp_path / "baitrage.png"
    photo.write_bytes(TINY_PNG)
    plan = SlidePlan(
        slides=[
            Slide(
                candidate_id="conclusion",
                free=True,
                title="Conclusion",
                bullets=["The strategy in action"],
                image="doc/baitrage.png",
            )
        ]
    )
    written = build(plan, tmp_path / "deck.pptx", images={"doc/baitrage.png": photo})
    slide = Presentation(str(written)).slides[0]

    picture = next(shape for shape in slide.shapes if hasattr(shape, "image"))
    assert picture.image.blob == TINY_PNG
    assert "The strategy in action" in _canvas_text(slide)
    assert "![" not in _canvas_text(slide)


def test_speaker_notes_stay_off_canvas_but_land_in_the_deck(tmp_path):
    plan = SlidePlan(
        slides=[
            Slide(
                candidate_id="auth",
                title="Auth",
                bullets=["Signed tokens"],
                speaker_notes="Mention the migration sequence.",
            )
        ]
    )
    slide = Presentation(str(build(plan, tmp_path / "notes.pptx"))).slides[0]
    assert "migration sequence" not in _canvas_text(slide)
    assert "migration sequence" in slide.notes_slide.notes_text_frame.text


def test_an_image_slide_with_no_resolved_path_fails_honestly(tmp_path):
    plan = SlidePlan(
        slides=[Slide(candidate_id="conclusion", free=True, title="Conclusion", image="doc/gone.png")]
    )
    with pytest.raises(InvalidInput, match="gone.png"):
        build(plan, tmp_path / "deck.pptx", images={})


def test_a_decompression_bomb_image_fails_honestly_not_with_a_raw_pillow_error(tmp_path):
    # A PNG header can declare dimensions Pillow refuses to decode without any real pixel
    # data - Image.open() raises DecompressionBombError on the header alone. That error
    # doesn't subclass OSError/ValueError, so it must be caught explicitly or it surfaces
    # as an unhandled exception instead of the same InvalidInput every other bad-image
    # path raises.
    def png_chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    huge_header = struct.pack(">IIBBBBB", 50_000, 50_000, 8, 6, 0, 0, 0)
    bomb = (
        bytes([137, 80, 78, 71, 13, 10, 26, 10])
        + png_chunk(b"IHDR", huge_header)
        + png_chunk(b"IEND", b"")
    )
    photo = tmp_path / "bomb.png"
    photo.write_bytes(bomb)

    plan = SlidePlan(
        slides=[Slide(candidate_id="conclusion", free=True, title="Conclusion", image="doc/bomb.png")]
    )
    with pytest.raises(InvalidInput, match="bomb.png"):
        build(plan, tmp_path / "deck.pptx", images={"doc/bomb.png": photo})
