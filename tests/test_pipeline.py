import threading
from datetime import UTC, datetime
from pathlib import Path

import pytest
from pptx import Presentation

from standup.core import index as index_module
from standup.core import pipeline
from standup.core.gather import candidates
from standup.core.index import docs
from standup.core.models import Scope, Slide
from standup.core.selection import choose
from standup.errors import InvalidInput, NotFound
from tests.conftest import TINY_PNG, pdf_saying


def slides_for(deck) -> list[Slide]:
    return [
        Slide(candidate_id=entry.candidate.id, title="A slide", bullets=["one"])
        for entry in deck.selection.chosen
    ]


def test_selecting_writes_a_deck_you_can_read_back(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup tomorrow", Scope(slide_budget=2))

    assert (store.paths(project.id).deck / pipeline.SELECTION_FILE).is_file()
    assert pipeline.read(store, project.id) == deck
    assert len(deck.selection.chosen) == 2
    assert deck.slides is None
    assert all(entry.signals for entry in deck.selection.chosen)


def test_sources_reads_the_real_file_a_candidate_id_points_at(store, project, index):
    """The index stores no content for a repo file, so a slide could only restate a commit message.
    The file is on disk the whole time; `sources` is what finally reaches it."""
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    ids = [entry.candidate.id for entry in deck.selection.chosen]

    found = pipeline.sources(store, project.id, ids)
    assert any("class Router" in text for text in found.values())
    assert all(len(text) <= pipeline.SOURCE_CHARS for text in found.values())


def test_sources_stays_quiet_about_a_file_that_is_no_longer_there(store, project, index):
    """A deleted file is still real work, and its commits still describe it."""
    assert pipeline.sources(store, project.id, ["nonexistent-resource/gone.py"]) == {}


def test_selecting_chooses_what_choose_would_have(store, project, index):
    scope = Scope(slide_budget=2)
    deck = pipeline.select(store, project.id, index, "standup", scope)
    direct = choose(index, scope, candidates(index, scope), request="standup")
    assert [e.candidate.id for e in deck.selection.chosen] == [
        e.candidate.id for e in direct.chosen
    ]


def test_a_second_select_keeps_a_slide_whose_candidate_is_still_chosen(store, project, index):
    """Used to delete every slide outright on any re-select, which is what made an ordinary
    follow-up wipe a finished deck - "the slides keep resetting"."""
    first = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(first))

    replaced = pipeline.select(store, project.id, index, "something else", Scope(slide_budget=2))
    assert replaced.selection.request == "something else"
    assert pipeline.read(store, project.id).slides == slides_for(first)


def test_a_second_select_drops_a_slide_whose_candidate_is_no_longer_chosen(store, project, index):
    """A window that excludes every prior candidate is a real change of scope, and a slide with no
    candidate behind it has genuinely lost the evidence it was written from."""
    first = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(first))

    future = Scope(slide_budget=2, since=datetime(2099, 1, 1, tzinfo=UTC))
    pipeline.select(store, project.id, index, "next year", future)
    assert pipeline.read(store, project.id).slides is None


def test_a_second_select_does_not_undo_a_keep_established_interleave(store, project, index):
    """Same bug shape as `write`'s reset above, left unpatched in `_surviving`: it still rebuilt the
    order as every evidence slide then every free one, so a `keep`-pinned free slide (a title screen)
    got thrown to the back on the very next re-select, even when nothing it depended on changed."""
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))
    pipeline.write(store, project.id, index, [Slide(candidate_id="title", free=True, title="Standup")])
    order = ["title", *[e.candidate.id for e in deck.selection.chosen]]
    pipeline.edit(store, project.id, order)

    reselected = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    assert [s.candidate_id for s in reselected.slides] == order


def test_an_edit_is_obeyed_literally(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    backwards = [entry.candidate.id for entry in reversed(deck.selection.chosen)]

    revised = pipeline.edit(store, project.id, backwards)
    assert [e.candidate.id for e in revised.selection.chosen] == backwards
    assert pipeline.read(store, project.id) == revised

    kept = backwards[:1]
    trimmed = pipeline.edit(store, project.id, kept)
    assert [e.candidate.id for e in trimmed.selection.chosen] == kept
    assert backwards[1] in {e.candidate.id for e in trimmed.selection.cut}


def test_an_edit_blocks_while_the_projects_lock_is_held(store, project, index):
    """select/edit/write/forget share one per-project lock — a worker thread and the event loop
    both reach them, and only a real lock (not a point-in-time busy check) is safe across that."""
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))

    with pipeline._LOCKS[project.id]:
        thread = threading.Thread(target=pipeline.edit, args=(store, project.id, []))
        thread.start()
        thread.join(timeout=0.3)
        assert thread.is_alive()

    thread.join(timeout=2)
    assert not thread.is_alive()


async def test_the_merged_index_is_cached_until_a_resource_actually_changes(store, project, repo, monkeypatch):
    calls = []
    real_merged = index_module.merged

    def spy(parts):
        calls.append(1)
        return real_merged(parts)

    monkeypatch.setattr(index_module, "merged", spy)

    first = await pipeline.indexed(store, project.id)
    second = await pipeline.indexed(store, project.id)
    assert len(calls) == 1
    assert first == second

    (repo / "src" / "new_file.py").write_text("def added():\n    pass\n")
    third = await pipeline.indexed(store, project.id)
    assert len(calls) == 2
    assert third != first


def test_read_blocks_while_the_projects_lock_is_held(store, project, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))

    with pipeline._LOCKS[project.id]:
        thread = threading.Thread(target=pipeline.read, args=(store, project.id))
        thread.start()
        thread.join(timeout=0.3)
        assert thread.is_alive()

    thread.join(timeout=2)
    assert not thread.is_alive()


def test_naming_the_same_item_twice_asks_for_it_once(store, project, index):
    """Repeats used to survive as duplicate entries, then render as one block copied per slide."""
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    twice = deck.selection.chosen[0].candidate.id

    revised = pipeline.edit(store, project.id, [twice, twice, twice])
    assert [entry.candidate.id for entry in revised.selection.chosen] == [twice]

    written = pipeline.write(store, project.id, index, slides_for(revised))
    assert [slide.candidate_id for slide in written.slides] == [twice]


def test_editing_to_something_that_was_never_a_candidate_is_refused(store, project, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    with pytest.raises(NotFound):
        pipeline.edit(store, project.id, ["src/imaginary.py"])


def test_writing_slides_that_match_the_selection_is_accepted(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    written = pipeline.write(store, project.id, index, slides_for(deck))

    assert (store.paths(project.id).deck / pipeline.PLAN_FILE).is_file()
    assert [s.candidate_id for s in written.slides] == [
        e.candidate.id for e in deck.selection.chosen
    ]


def test_writing_one_selected_slide_leaves_the_rest_unwritten(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    first = slides_for(deck)[:1]

    written = pipeline.write(store, project.id, index, first)

    assert written.slides == first
    assert pipeline.read(store, project.id).slides == first


def test_writing_remaining_slides_merges_them_in_selection_order(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    drafted = slides_for(deck)
    pipeline.write(store, project.id, index, drafted[1:])

    written = pipeline.write(store, project.id, index, drafted[:1])

    assert written.slides == drafted


def test_a_partial_slide_still_must_be_grounded(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    bad = slides_for(deck)[:1]
    bad[0].bullets = ["Rewrote src/imaginary.py"]

    with pytest.raises(InvalidInput, match="no file 'src/imaginary.py'"):
        pipeline.write(store, project.id, index, bad)
    assert not (store.paths(project.id).deck / pipeline.PLAN_FILE).is_file()


def test_rendering_a_partial_plan_is_refused(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck)[:1])

    with pytest.raises(InvalidInput, match="not complete"):
        pipeline.render(store, project.id)


def test_writing_a_slide_for_something_not_chosen_is_refused(store, project, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=1))
    bogus = Slide(candidate_id="src/imaginary.py", title="Ghost", bullets=["x"])

    with pytest.raises(NotFound, match="src/imaginary.py"):
        pipeline.write(store, project.id, index, [bogus])
    assert not (store.paths(project.id).deck / pipeline.PLAN_FILE).is_file()


def test_a_free_slide_needs_no_candidate(store, project, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    free = Slide(candidate_id="title", free=True, title="Standup", bullets=[])

    written = pipeline.write(store, project.id, index, [free])
    assert written.slides == [free]


def test_a_deck_cannot_be_completed_with_only_free_slides(store, project, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    slides = [
        Slide(candidate_id="intro", free=True, title="Intro"),
        Slide(candidate_id="wrap", free=True, title="Wrap"),
    ]

    pipeline.write(store, project.id, index, slides)
    with pytest.raises(InvalidInput, match="evidence slides"):
        pipeline.render(store, project.id)


def test_a_new_free_slide_lands_after_the_evidence_slides(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))
    free = Slide(candidate_id="title", free=True, title="Standup")

    written = pipeline.write(store, project.id, index, [free])
    assert written.slides[-1] == free
    assert [s.candidate_id for s in written.slides[:-1]] == [
        e.candidate.id for e in deck.selection.chosen
    ]


def test_keep_can_interleave_a_free_slide_among_evidence_slides(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))
    pipeline.write(store, project.id, index, [Slide(candidate_id="title", free=True, title="Standup")])

    order = ["title", *[e.candidate.id for e in deck.selection.chosen]]
    reordered = pipeline.edit(store, project.id, order)
    assert [s.candidate_id for s in reordered.slides] == order
    assert pipeline.render(store, project.id).is_file()


def test_writing_a_different_slide_does_not_undo_a_keep_established_interleave(store, project, index):
    """Real bug, reported as "you reshifted the introduction slide back" after nothing more than
    adding an image to a different slide. `write` used to rebuild the whole order as every evidence
    slide then every free one, discarding wherever `keep` had just placed a free slide among them."""
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))
    pipeline.write(store, project.id, index, [Slide(candidate_id="title", free=True, title="Standup")])
    order = ["title", *[e.candidate.id for e in deck.selection.chosen]]
    pipeline.edit(store, project.id, order)

    evidence_id = deck.selection.chosen[0].candidate.id
    written = pipeline.write(
        store, project.id, index,
        [Slide(candidate_id=evidence_id, title="Renamed", bullets=["still true"])],
    )
    assert [s.candidate_id for s in written.slides] == order


def test_keep_without_the_free_id_drops_it(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))
    pipeline.write(store, project.id, index, [Slide(candidate_id="title", free=True, title="Standup")])

    kept = [e.candidate.id for e in deck.selection.chosen]
    reordered = pipeline.edit(store, project.id, kept)
    assert "title" not in [s.candidate_id for s in reordered.slides]


async def test_rendering_an_image_slide_embeds_the_attached_files_actual_bytes(store, project, monkeypatch):
    """The resource an image id names is a lone attached file - its `location` is the file itself,
    not a folder to walk into."""
    monkeypatch.setattr(docs, "_DESCRIBED", {})
    monkeypatch.setattr(docs, "describe_image_sync", lambda data, media_type, *, prompt: "A chart")
    resource = store.attach_file(project.id, "chart.png", TINY_PNG)
    index = await pipeline.indexed(store, project.id)

    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=1))
    pipeline.write(store, project.id, index, slides_for(deck))
    image = Slide(candidate_id="wrap", free=True, title="Wrap", image=f"{resource.id}/chart.png")
    pipeline.write(store, project.id, index, [image])

    written = pipeline.render(store, project.id)
    reopened = Presentation(str(written))
    picture = next(shape for shape in reopened.slides[-1].shapes if hasattr(shape, "image"))
    assert picture.image.blob == TINY_PNG


def test_a_free_flag_on_a_real_candidate_is_refused(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    real_id = deck.selection.chosen[0].candidate.id
    sneaky = Slide(candidate_id=real_id, free=True, title="Sneaky", bullets=["src/imaginary.py"])

    with pytest.raises(InvalidInput, match="cannot be free"):
        pipeline.write(store, project.id, index, [sneaky])


def test_a_free_flag_on_a_cut_candidate_is_also_refused(store, project, index):
    """Not just chosen — a free slide squatting on a cut candidate's id would corrupt it if
    that candidate were later restored with `keep`."""
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=1))
    assert deck.selection.cut
    cut_id = deck.selection.cut[0].candidate.id
    sneaky = Slide(candidate_id=cut_id, free=True, title="Sneaky")

    with pytest.raises(InvalidInput, match="cannot be free"):
        pipeline.write(store, project.id, index, [sneaky])


def test_rendering_an_empty_plan_is_refused_rather_than_producing_an_empty_deck(store, project, index):
    """A SlidePlan(slides=[]) is truthy as a Pydantic model — render() must check its content."""
    future = datetime(2099, 1, 1, tzinfo=UTC)
    empty = pipeline.select(store, project.id, index, "nothing here", Scope(since=future, slide_budget=1))
    assert empty.selection.chosen == []

    pipeline.write(store, project.id, index, [])
    with pytest.raises(InvalidInput, match="No slides"):
        pipeline.render(store, project.id)


def test_a_slide_naming_something_that_does_not_exist_is_rejected(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    drafted = slides_for(deck)
    drafted[0].bullets = ["Rewrote src/imaginary.py"]

    with pytest.raises(InvalidInput, match="no file 'src/imaginary.py'"):
        pipeline.write(store, project.id, index, drafted)
    assert not (store.paths(project.id).deck / pipeline.PLAN_FILE).is_file()


async def test_detaching_a_resource_takes_the_deck_it_produced_with_it(store, project):
    """Left behind, its ids reach the agent as evidence and read as context that still exists."""
    paper = store.attach_file(project.id, "notes.pdf", pdf_saying("Router notes"))
    both = await pipeline.indexed(store, project.id)
    deck = pipeline.select(store, project.id, both, "everything", Scope(slide_budget=5))
    pipeline.write(store, project.id, both, slides_for(deck))
    assert any(entry.candidate.id.startswith(paper.id) for entry in deck.selection.chosen)

    store.detach(project.id, paper.id)
    pipeline.forget(store, project.id, paper.id)

    left = pipeline.read(store, project.id)
    named = [entry.candidate.id for entry in [*left.selection.chosen, *left.selection.cut]]
    assert named and not any(item.startswith(paper.id) for item in named)
    assert not any(slide.candidate_id.startswith(paper.id) for slide in left.slides or [])
    assert not (store.paths(project.id).deck / pipeline.DECK_FILE).is_file()


def test_detaching_the_last_resource_leaves_no_deck_at_all(store, project, resource_id, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))

    store.detach(project.id, resource_id)
    pipeline.forget(store, project.id, resource_id)

    with pytest.raises(NotFound):
        pipeline.read(store, project.id)


def test_rewriting_one_slide_leaves_the_others_untouched(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))
    before = pipeline.read(store, project.id).slides

    reworded = Slide(candidate_id=before[0].candidate_id, title="Reworded", bullets=["new"])
    after = pipeline.write(store, project.id, index, [reworded])

    assert after.slides[0] == reworded
    assert after.slides[1:] == before[1:]


def test_reordering_carries_the_slides_with_it(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))

    backwards = [entry.candidate.id for entry in reversed(deck.selection.chosen)]
    reordered = pipeline.edit(store, project.id, backwards)
    assert [s.candidate_id for s in reordered.slides] == backwards


def test_promoting_a_cut_item_leaves_it_needing_a_slide_without_losing_what_was_already_written(
    store, project, index
):
    """A keep that brings back a never-written cut candidate used to wipe every slide already
    written for the rest of the deck, not just leave the new one unwritten."""
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=1))
    written = pipeline.write(store, project.id, index, slides_for(deck))
    already_written = written.slides[0].candidate_id

    everything = [e.candidate.id for e in [*deck.selection.chosen, *deck.selection.cut]]
    promoted = pipeline.edit(store, project.id, everything)

    assert [e.candidate.id for e in promoted.selection.chosen] == everything
    assert [s.candidate_id for s in promoted.slides] == [already_written]


def test_rendering_needs_slides_first(store, project, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    with pytest.raises(InvalidInput, match="No slides"):
        pipeline.render(store, project.id)


def test_rendering_writes_a_deck_file(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    pipeline.write(store, project.id, index, slides_for(deck))
    assert pipeline.render(store, project.id).is_file()


def test_a_project_with_no_deck_yet_is_not_found(store, project):
    with pytest.raises(NotFound):
        pipeline.read(store, project.id)


def test_a_corrupt_selection_file_is_a_clean_error_not_a_crash(store, project, index):
    pipeline.select(store, project.id, index, "standup", Scope(slide_budget=1))
    (store.paths(project.id).deck / pipeline.SELECTION_FILE).write_text("not json", encoding="utf-8")

    with pytest.raises(InvalidInput, match="corrupt"):
        pipeline.read(store, project.id)


def test_a_corrupt_plan_file_is_a_clean_error_not_a_crash(store, project, index):
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=1))
    pipeline.write(store, project.id, index, slides_for(deck))
    (store.paths(project.id).deck / pipeline.PLAN_FILE).write_text("{not json", encoding="utf-8")

    with pytest.raises(InvalidInput, match="corrupt"):
        pipeline.read(store, project.id)


def test_a_crash_mid_write_never_corrupts_the_file_a_reader_opens(store, project, index, monkeypatch):
    first = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=2))
    good = (store.paths(project.id).deck / pipeline.SELECTION_FILE).read_text(encoding="utf-8")

    real_replace = Path.replace

    def crash_before_replace(self, target):
        raise OSError("simulated crash")

    monkeypatch.setattr(Path, "replace", crash_before_replace)
    with pytest.raises(OSError, match="simulated crash"):
        pipeline.select(store, project.id, index, "standup again", Scope(slide_budget=1))

    monkeypatch.setattr(Path, "replace", real_replace)
    assert (store.paths(project.id).deck / pipeline.SELECTION_FILE).read_text(encoding="utf-8") == good
    assert pipeline.read(store, project.id) == first
