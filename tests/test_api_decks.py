import io

import httpx
from pptx import Presentation

from standup.api import jobs, routes
from standup.api import projects as projects_api
from standup.api.app import app
from standup.core import pipeline
from standup.core.models import Scope, Slide
from tests.conftest import Stuck
from tests.conftest import settled as await_settled


async def test_editing_the_selection_is_refused_while_a_turn_is_running(store, monkeypatch, project):
    monkeypatch.setattr(projects_api, "get_store", lambda: store)
    stuck = Stuck()
    monkeypatch.setattr(routes, "get_client", lambda: stuck)

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        turn = await client.post(f"/projects/{project.id}/chat", json={"content": "go"})
        assert turn.status_code == 202

        refused = await client.put(f"/projects/{project.id}/deck/selection", json={"keep": []})
        assert refused.status_code == 409

        stuck.released.set()
        job = await await_settled(client, turn.json()["id"])
        assert job["state"] == "done"


async def test_downloading_the_deck_returns_real_openable_pptx_bytes(store, monkeypatch, project, index):
    """The GUI used to have no way to reach this endpoint at all - proves what it now points at is a
    real, PowerPoint-openable file over HTTP, not just a 200 and a promise."""
    monkeypatch.setattr(projects_api, "get_store", lambda: store)
    deck = pipeline.select(store, project.id, index, "standup", Scope(slide_budget=1))
    pipeline.write(
        store, project.id, index,
        [Slide(candidate_id=deck.selection.chosen[0].candidate.id, title="Router", bullets=["a"])],
    )

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.get(f"/projects/{project.id}/deck/file")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    opened = Presentation(io.BytesIO(response.content))
    canvas = "\n".join(shape.text for shape in opened.slides[0].shapes if shape.has_text_frame)
    assert "Router" in canvas


async def test_editing_the_selection_is_refused_while_indexing_is_in_progress(store, monkeypatch, project):
    """busy() must see indexing too, not only a chat turn — indexing runs unlocked in _LOOSE."""
    monkeypatch.setattr(projects_api, "get_store", lambda: store)

    async with (
        jobs.indexing_lock(project.id),
        httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client,
    ):
        refused = await client.put(f"/projects/{project.id}/deck/selection", json={"keep": []})
        assert refused.status_code == 409
